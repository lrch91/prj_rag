"""PPTX 解析器"""

from .base import BaseParser, ParsedPage


class PptxParser(BaseParser):
    def parse(self, file_path: str) -> list[ParsedPage]:
        from pptx import Presentation

        prs = Presentation(file_path)
        pages: list[ParsedPage] = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
            if texts:
                pages.append(ParsedPage(page_number=i, text="\n".join(texts)))
        return pages
