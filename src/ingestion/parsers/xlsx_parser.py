"""XLSX 解析器"""

from .base import BaseParser, ParsedPage


class XlsxParser(BaseParser):
    def parse(self, file_path: str) -> list[ParsedPage]:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        parts: list[str] = []
        for name in wb.sheetnames:
            ws = wb[name]
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"Sheet: {name}\n" + "\n".join(rows))
        wb.close()
        text = "\n\n".join(parts)
        if not text.strip():
            return []
        return [ParsedPage(page_number=1, text=text)]
